"""
One loader per file type. Every loader returns the same format, a list of pages:

    [{"text": "...", "course": "nlp", "file": "nlp_notes.pdf", "page": 1}, ...]

Because the format is identical, chunking and indexing never need to know
whether a page came from a PDF, a slide deck or a text file.
"""

import re
from collections import Counter
from pathlib import Path

from pypdf import PdfReader
from pptx import Presentation

import config


def _page(text, course, path, page):
    """Build one page record in the shared format."""
    return {"text": text.strip(), "course": course, "file": Path(path).name, "page": page}


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def _join_lone_bullets(text):
    """pypdf often puts a bullet on its own line ("-\\nQuizzes: 15%").
    Join it back onto the next line so the bullet and its text stay together."""
    return re.sub(r"(?m)^([-•*])\s*\n", r"\1 ", text)


def _reflow_wrapped_lines(text, min_wrap_len=60):
    """pypdf keeps the PDF's visual line wraps, so one sentence arrives split
    over several lines. Join a line onto the next one when it looks like a wrap:
    long (near the full page width) and not ending in sentence punctuation.
    Short lines (headings) and lines ending in "." stay separate, so the text
    keeps one line per heading / bullet / paragraph, which chunking relies on."""
    out = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        prev = out[-1] if out else ""
        if len(prev) >= min_wrap_len and not prev.endswith((".", "!", "?", ":")):
            out[-1] = f"{prev} {line}"
        else:
            out.append(line)
    return "\n".join(out)


def _normalize_line(line):
    """Turn "Page 3" and "Page 4" into the same key ("Page #"), so footers
    that only differ by their page number are still recognised as repeats."""
    return re.sub(r"\d+", "#", line.strip())


def strip_repeated_lines(page_texts):
    """Remove header/footer lines that repeat on most pages of one file.

    A line is treated as boilerplate if (after replacing digits with #) it
    appears on at least REPEATED_LINE_THRESHOLD of the pages. Left in, the
    same footer would land in every chunk and make all chunks look alike to
    the embedding model.

    Returns (cleaned_texts, removed_lines).
    """
    # With only 1-2 pages, "repeats on most pages" is meaningless: any shared
    # line would be removed. Leave short files alone.
    if len(page_texts) < 3:
        return page_texts, set()

    # Count each normalised line once per page (a set), so a line repeated
    # several times on ONE page doesn't look like a footer.
    counts = Counter()
    for text in page_texts:
        counts.update({_normalize_line(l) for l in text.splitlines() if l.strip()})

    min_pages = config.REPEATED_LINE_THRESHOLD * len(page_texts)
    repeated = {line for line, n in counts.items() if n >= min_pages}

    cleaned = []
    for text in page_texts:
        kept = [l for l in text.splitlines() if _normalize_line(l) not in repeated]
        cleaned.append("\n".join(kept))
    return cleaned, repeated


def load_pdf(path, course):
    reader = PdfReader(path)
    raw = [_join_lone_bullets(p.extract_text() or "") for p in reader.pages]
    # Strip footers BEFORE reflowing, while they are still separate lines.
    cleaned, _ = strip_repeated_lines(raw)
    cleaned = [_reflow_wrapped_lines(t) for t in cleaned]

    # Page numbers start at 1 to match what a student sees in the PDF viewer.
    pages = [_page(t, course, path, i + 1) for i, t in enumerate(cleaned) if t.strip()]
    if not pages:
        print(f"  WARNING: no extractable text in {Path(path).name} "
              "(scanned PDF? it needs OCR first). Skipping.")
    return pages


# ---------------------------------------------------------------------------
# PowerPoint
# ---------------------------------------------------------------------------
def load_pptx(path, course):
    prs = Presentation(path)
    pages = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None else ""
        # Compare by shape_id: python-pptx creates a new Python object on every
        # access, so `shape is title_shape` would never be True.
        title_id = title_shape.shape_id if title_shape is not None else None

        # Collect text from every other shape that holds text (body boxes, etc.).
        body = []
        for shape in slide.shapes:
            if shape.shape_id == title_id or not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if text:
                body.append(text)

        # Put the title first: slide bodies are often terse bullets that only
        # make sense with their title ("Pros / Cons" of WHAT?).
        text = "\n".join([title] + body) if title else "\n".join(body)
        if text.strip():
            pages.append(_page(text, course, path, slide_num))

    if not pages:
        print(f"  WARNING: no extractable text in {Path(path).name} "
              "(image-only slides?). Skipping.")
    return pages


# ---------------------------------------------------------------------------
# Plain text / Markdown
# ---------------------------------------------------------------------------
def load_text(path, course):
    # errors="replace" so one odd character doesn't crash the whole indexing run.
    text = Path(path).read_text(encoding="utf-8", errors="replace")
    if not text.strip():
        print(f"  WARNING: {Path(path).name} is empty. Skipping.")
        return []
    # Text files have no pages, so the whole file is "page 1".
    return [_page(text, course, path, 1)]


# ---------------------------------------------------------------------------
# Dispatcher
# ---------------------------------------------------------------------------
LOADERS = {".pdf": load_pdf, ".pptx": load_pptx, ".txt": load_text, ".md": load_text}


def load_file(path, course):
    """Load any supported file. Unsupported types return [] with a warning."""
    loader = LOADERS.get(Path(path).suffix.lower())
    if loader is None:
        print(f"  WARNING: unsupported file type, skipping: {Path(path).name}")
        return []
    try:
        return loader(path, course)
    except Exception as e:
        # A corrupt file shouldn't stop the other files from being indexed.
        print(f"  WARNING: could not read {Path(path).name}: {e}. Skipping.")
        return []
